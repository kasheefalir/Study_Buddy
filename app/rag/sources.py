"""Deterministic extraction; original files and web responses remain untouched."""
from pathlib import Path
import zipfile
from xml.etree import ElementTree
from bs4 import BeautifulSoup
from pptx import Presentation


def slide_text(shapes):
    for shape in shapes:
        if hasattr(shape, "shapes"):
            yield from slide_text(shape.shapes)
        if shape.has_text_frame:
            yield shape.text_frame.text
        if shape.has_table:
            for row in shape.table.rows:
                yield " | ".join(cell.text for cell in row.cells)


def extract_pptx(path: Path):
    with zipfile.ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Expanded PowerPoint exceeds the 100 MB extraction limit.")
    presentation = Presentation(path)
    passages = []
    for number, slide in enumerate(presentation.slides, 1):
        text = "\n".join(slide_text(slide.shapes)).strip()
        if text:
            passages.append({"location": f"Slide {number}", "text": text})
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame and frame.text.strip():
                passages.append({"location": f"Slide {number} · speaker notes", "text": frame.text})
    return passages


def extract_html(raw: bytes):
    soup = BeautifulSoup(raw, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    for element in soup(["script", "style", "noscript", "template"]):
        element.decompose()
    # Keep navigation and sidebars too, so later selection can recover material
    # that a brittle main-content heuristic would otherwise discard.
    text = soup.get_text("\n", strip=True)
    return title, [{"location": "Page text", "text": text}]


def extract_docx(path: Path):
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    passages = []
    with zipfile.ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Expanded Word document exceeds the 100 MB extraction limit.")
        document = ElementTree.fromstring(archive.read("word/document.xml"))
        # Paragraphs include table cells in document order. No macros or external
        # relationships are evaluated, and original files remain untouched.
        for number, paragraph in enumerate(document.iter(namespace + "p"), 1):
            text = "".join(node.text or "" for node in paragraph.iter(namespace + "t")).strip()
            if text:
                passages.append({"location": f"Paragraph {number}", "text": text})
    return passages


def chunks(passages):
    result = []
    for passage in passages:
        text = passage["text"].strip()
        if passage.get("variant") or passage.get("page"):
            start = 0
            while start < len(text):
                end = min(start + 900, len(text))
                if end < len(text):
                    boundary = max(text.rfind(" ", start + 450, end), text.rfind("\n", start + 450, end))
                    if boundary > start:
                        end = boundary
                result.append({**passage, "id": len(result) + 1, "text": text[start:end]})
                if end == len(text):
                    break
                overlap = max(start + 1, end - 180)
                while overlap < end and not text[overlap].isspace():
                    overlap += 1
                start = overlap + 1
            continue
        for start in range(0, len(text), 700):
            excerpt = text[start:start + 900]
            if excerpt.strip():
                result.append({**passage, "id": len(result) + 1, "text": excerpt})
    return result
