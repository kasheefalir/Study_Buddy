import asyncio
import json
import socket

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.rag import library, web_import
from app.rag.sources import extract_pptx, extract_html
from app.agents.chat_agent import ChatAgent
from app.storage.conversation_store import ConversationStore
from app.llm.usage import record_usage


def deck(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Regularization"
    slide.placeholders[1].text = "Regularization reduces overfitting by penalizing model complexity."
    slide.notes_slide.notes_text_frame.text = "L2 regularization penalizes squared coefficient values."
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Method"
    table.cell(0, 1).text = "Penalty"
    table.cell(1, 0).text = "L1"
    table.cell(1, 1).text = "Absolute values"
    presentation.save(path)


def test_powerpoint_sources_and_notes_retained(tmp_path):
    identifier = "a" * 32
    source = tmp_path / (identifier + ".pptx")
    deck(source)
    original = source.read_bytes()
    passages = extract_pptx(source)
    assert any("Absolute values" in p["text"] for p in passages)
    assert any("squared" in p["text"] and "notes" in p["location"] for p in passages)
    item = {"id": identifier, "kind": "pptx", "name": "Week5.pptx", "filename": source.name}
    assert library.process(tmp_path, item)["status"] == "Ready"
    assert source.read_bytes() == original
    assert (tmp_path / identifier / "extracted.json").exists()
    library.process(tmp_path, item)
    assert list((tmp_path / identifier / "versions").glob("*/extracted.json"))


def test_html_preserves_secondary_text_but_not_scripts():
    _, passages = extract_html(b"<title>Lesson</title><nav>Extra definitions</nav><main>Course text</main><script>evil()</script>")
    assert "Extra definitions" in passages[0]["text"]
    assert "Course text" in passages[0]["text"]
    assert "evil()" not in passages[0]["text"]


def test_web_snapshot_and_reimport_versions(tmp_path, monkeypatch):
    raw = b"<title>Study</title><p>Photosynthesis uses light.</p>"
    monkeypatch.setattr(library, "fetch_page", lambda url: (raw, {"content_type": "text/html", "final_url": url}))
    item = {"id": "b" * 32, "kind": "url", "name": "https://example.com", "url": "https://example.com"}
    result = library.process(tmp_path, item)
    assert result["status"] == "Ready"
    archive = tmp_path / item["id"]
    assert (archive / "original.bin").read_bytes() == raw
    library.process(tmp_path, item)
    assert list((archive / "versions").glob("*/original.bin"))


@pytest.mark.parametrize("ip", ["127.0.0.1", "10.0.0.1", "169.254.169.254", "::1"])
def test_private_network_fetch_is_rejected(monkeypatch, ip):
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **kw: [(2, 1, 6, "", (ip, 80))])
    with pytest.raises(ValueError, match="Private"):
        web_import.resolve_public("http://example.com")


def test_grounding_selection_and_source_persistence(tmp_path):
    item = {"id": "c" * 32, "kind": "txt", "name": "Notes", "filename": "notes.txt"}
    (tmp_path / "notes.txt").write_text("Regularization penalizes model complexity to reduce overfitting.")
    library.process(tmp_path, item)

    class Model:
        def __init__(self):
            self.calls = []

        async def chat(self, messages):
            self.calls.append(messages)
            record_usage({"prompt_eval_count": 10, "eval_count": 2}, "test")
            return "[1]" if len(self.calls) == 1 else "It reduces overfitting [S1]."

    model = Model()
    store = ConversationStore(tmp_path / "conversations")
    identifier, reply = asyncio.run(ChatAgent(model, store).reply("What do my notes say about regularization?", document_id=item["id"], materials_root=tmp_path))
    assert "[S1]" in reply
    assert "Reviewed session materials" in model.calls[1][-2]["content"]
    assert store.document(identifier) == item["id"]
    assert store.load(identifier)[-1]["sources"][0]["location"] == "Document text"
    assert list((tmp_path / item["id"] / "retrievals").glob("*.json"))
    usage = store.load(identifier)[-1]["usage"]
    assert usage["input_tokens"] == 20
    assert usage["output_tokens"] == 4
    assert [call["stage"] for call in usage["calls"]] == ["study planning", "answer"]
    assert reply == "It reduces overfitting [S1]."
    assert "selected original passages are the source of truth" in model.calls[1][0]["content"]
    assert "For general tutoring" in model.calls[1][0]["content"]


def test_bad_ai_selection_falls_back_without_erasing_text(tmp_path):
    item = {"id": "d" * 32, "kind": "txt", "name": "Notes", "filename": "notes.txt"}
    (tmp_path / "notes.txt").write_text("A source fact.")
    library.process(tmp_path, item)

    class Model:
        async def chat(self, messages):
            return "[999]"
    sources = asyncio.run(library.ground(tmp_path, item["id"], "fact", Model()))
    assert sources[0]["text"] == "A source fact."
    audit = next((tmp_path / item["id"] / "retrievals").glob("*.json"))
    assert json.loads(audit.read_text())["selection"] == "lexical fallback"


def test_overview_covers_end_of_document(tmp_path):
    item = {"id": "e" * 32, "name": "Course", "status": "Ready"}
    library.write_json(tmp_path / (item["id"] + ".json"), item)
    indexed = [{"id": i, "location": f"Slide {i}", "text": f"Topic {i}"} for i in range(1, 31)]
    library.write_json(tmp_path / item["id"] / "chunks.json", indexed)

    class Model:
        async def chat(self, messages):
            return '{"mode":"overview","query":""}'

    sources = asyncio.run(library.ground(tmp_path, item["id"], "Summarize", Model()))
    assert len(sources) == 12
    assert sources[0]["location"] == "Slide 1"
    assert sources[-1]["location"] == "Slide 30"


def test_planned_followup_search_uses_resolved_topic(tmp_path):
    item = {"id": "f" * 32, "name": "Course", "status": "Ready"}
    library.write_json(tmp_path / (item["id"] + ".json"), item)
    library.write_json(tmp_path / item["id"] / "chunks.json", [
        {"id": 1, "location": "Slide 1", "text": "Classification cutoff"},
        {"id": 2, "location": "Slide 2", "text": "Regularization shrinks coefficients"},
    ])

    class Model:
        async def chat(self, messages):
            return '{"mode":"search","query":"regularization coefficients"}'

    sources = asyncio.run(library.ground(tmp_path, item["id"], "Explain that simply", Model()))
    assert sources[0]["chunk_id"] == 2


def test_planned_search_uses_semantic_retrieval_for_text_materials(tmp_path, monkeypatch):
    item = {"id": "1" * 32, "kind": "txt", "name": "Notes", "status": "Ready"}
    library.write_json(tmp_path / (item["id"] + ".json"), item)
    chunks = [{"id": 1, "location": "Notes", "text": "Penalty terms discourage overly complex models."}]
    library.write_json(tmp_path / item["id"] / "chunks.json", chunks)
    called = []

    async def semantic(root, material, indexed, query, limit):
        called.append((material["kind"], query, limit))
        return indexed

    monkeypatch.setattr("app.rag.semantic.retrieve", semantic)

    class Model:
        async def chat(self, messages):
            return '{"mode":"search","query":"avoid memorizing examples"}'

    sources = asyncio.run(library.ground(tmp_path, item["id"], "Why regularize?", Model()))
    assert sources[0]["chunk_id"] == 1
    assert called == [("txt", "avoid memorizing examples", 6)]
